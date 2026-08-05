from __future__ import annotations

from pathlib import Path
from datetime import date

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "outputs" / "Unity3D_Web3_FPS_游戏设计需求文档_v1.0.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
NAVY = "17365D"
MUTED = "666666"
LIGHT_GRAY = "F2F4F7"
CALL_OUT = "E8EEF5"
WHITE = "FFFFFF"
GREEN = "E2F0D9"
AMBER = "FFF2CC"
RED = "FCE4D6"
BLACK = "000000"

FONT_LATIN = "Calibri"
FONT_CJK = "Microsoft YaHei"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D9DEE5", size=4) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:color"), color)


def set_table_geometry(table, widths_dxa: list[int], indent_dxa: int = 120) -> None:
    assert sum(widths_dxa) == 9360, widths_dxa
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl = table._tbl
    tbl_pr = tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(widths_dxa[idx] / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def set_run_font(run, size=None, color=None, bold=None, italic=None, name=FONT_LATIN) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT_CJK)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def style_paragraph_runs(paragraph, size=11, color=BLACK, bold=None) -> None:
    for run in paragraph.runs:
        set_run_font(run, size=size, color=color, bold=bold if bold is not None else run.bold)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = paragraph.add_run("第 ")
    set_run_font(r, size=9, color=MUTED)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    rr = paragraph.add_run()
    rr._r.append(fld_char1)
    rr._r.append(instr)
    rr._r.append(fld_char2)
    set_run_font(rr, size=9, color=MUTED)
    r2 = paragraph.add_run(" 页")
    set_run_font(r2, size=9, color=MUTED)


def add_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    rid = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rid)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(color)
    r_pr.append(underline)
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), FONT_LATIN)
    r_fonts.set(qn("w:hAnsi"), FONT_LATIN)
    r_fonts.set(qn("w:eastAsia"), FONT_CJK)
    r_pr.append(r_fonts)
    new_run.append(r_pr)
    text_el = OxmlElement("w:t")
    text_el.text = text
    new_run.append(text_el)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def add_custom_numbering(doc: Document) -> tuple[int, int]:
    numbering = doc.part.numbering_part.element
    existing_abstract_ids = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    existing_num_ids = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    next_abs = max(existing_abstract_ids or [0]) + 1
    next_num = max(existing_num_ids or [0]) + 1

    def make_num(abstract_id: int, num_id: int, fmt: str, text: str, font: str = FONT_LATIN):
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abstract_id))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "singleLevel")
        abstract.append(multi)
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        lvl.append(start)
        num_fmt = OxmlElement("w:numFmt")
        num_fmt.set(qn("w:val"), fmt)
        lvl.append(num_fmt)
        lvl_text = OxmlElement("w:lvlText")
        lvl_text.set(qn("w:val"), text)
        lvl.append(lvl_text)
        suff = OxmlElement("w:suff")
        suff.set(qn("w:val"), "tab")
        lvl.append(suff)
        p_pr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), "720")
        tabs.append(tab)
        p_pr.append(tabs)
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), "720")
        ind.set(qn("w:hanging"), "360")
        p_pr.append(ind)
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:after"), "160")
        spacing.set(qn("w:line"), "280")
        spacing.set(qn("w:lineRule"), "auto")
        p_pr.append(spacing)
        lvl.append(p_pr)
        r_pr = OxmlElement("w:rPr")
        r_fonts = OxmlElement("w:rFonts")
        r_fonts.set(qn("w:ascii"), font)
        r_fonts.set(qn("w:hAnsi"), font)
        r_fonts.set(qn("w:eastAsia"), FONT_CJK)
        r_pr.append(r_fonts)
        lvl.append(r_pr)
        abstract.append(lvl)
        numbering.append(abstract)

        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        aid = OxmlElement("w:abstractNumId")
        aid.set(qn("w:val"), str(abstract_id))
        num.append(aid)
        numbering.append(num)

    make_num(next_abs, next_num, "bullet", "•", FONT_LATIN)
    make_num(next_abs + 1, next_num + 1, "decimal", "%1.", FONT_LATIN)
    return next_num, next_num + 1


def add_list_item(doc, text: str, num_id: int, bold_prefix: str | None = None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.167
    p_pr = p._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_el = OxmlElement("w:numId")
    num_id_el.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num_id_el)
    p_pr.append(num_pr)
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        set_run_font(r, size=11, bold=True)
        rest = p.add_run(text[len(bold_prefix):])
        set_run_font(rest, size=11)
    else:
        r = p.add_run(text)
        set_run_font(r, size=11)
    return p


def add_heading(doc, text: str, level: int = 1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.add_run(text)
    return p


def add_body(doc, text: str, bold_prefix: str | None = None, italic=False):
    p = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2)
    else:
        r = p.add_run(text)
        set_run_font(r, italic=italic)
    return p


def add_callout(doc, label: str, text: str, fill: str = CALL_OUT):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.08)
    p.paragraph_format.right_indent = Inches(0.08)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(10)
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    borders = OxmlElement("w:pBdr")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "6")
        el.set(qn("w:color"), fill)
        borders.append(el)
    p_pr.append(borders)
    r1 = p.add_run(label + "：")
    set_run_font(r1, size=10.5, bold=True, color=NAVY)
    r2 = p.add_run(text)
    set_run_font(r2, size=10.5, color=BLACK)


def add_table(doc, headers: list[str], rows: list[list[str]], widths: list[int], alignments=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, text in enumerate(headers):
        cell = hdr.cells[i]
        set_cell_shading(cell, LIGHT_GRAY)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        set_run_font(r, size=9.5, bold=True, color=NAVY)
    for row_data in rows:
        row = table.add_row()
        for i, text in enumerate(row_data):
            cell = row.cells[i]
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.08
            if alignments:
                p.alignment = alignments[i]
            r = p.add_run(text)
            set_run_font(r, size=9.2)
    set_table_geometry(table, widths)
    set_table_borders(table)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def configure_styles(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = FONT_LATIN
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT_LATIN)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_LATIN)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CJK)
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    tokens = {
        1: (16, BLUE, 16, 8),
        2: (13, BLUE, 12, 6),
        3: (12, DARK_BLUE, 8, 4),
    }
    for level, (size, color, before, after) in tokens.items():
        style = doc.styles[f"Heading {level}"]
        style.font.name = FONT_LATIN
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT_LATIN)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_LATIN)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CJK)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header
    hp = header.paragraphs[0]
    hp.clear()
    hp.paragraph_format.space_after = Pt(0)
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = hp.add_run("WEB3 FPS · UNITY 3D 游戏侧需求")
    set_run_font(r, size=8.5, color=MUTED, bold=True)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.clear()
    add_page_number(fp)


def build_document():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_styles(doc)
    bullet_id, number_id = add_custom_numbering(doc)

    # First-page memo masthead — named title override: 24 pt navy.
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("游戏设计需求文档")
    set_run_font(r, size=24, color=NAVY, bold=True)
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(14)
    r = p2.add_run("Unity 3D Web3 FPS · 游戏侧与链上资产层分离")
    set_run_font(r, size=13, color=MUTED)

    meta = [
        ("文档版本", "v1.0"),
        ("文档类型", "游戏侧产品与集成需求（PRD / GDD 衔接稿）"),
        ("适用对象", "游戏策划、Unity 客户端、权威游戏服务器、后端与 QA"),
        ("基准日期", "2026-08-05"),
        ("技术基线", "Unity 3D + C#；实时对战链下；资产与结算能力经后端接入"),
    ]
    for label, value in meta:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        r1 = p.add_run(label + "：")
        set_run_font(r1, size=10.5, bold=True, color=NAVY)
        r2 = p.add_run(value)
        set_run_font(r2, size=10.5)

    add_callout(doc, "核心原则", "任何需要在实时对战帧内完成的行为都不得依赖区块链。Unity 客户端不持有私钥、不直连 RPC；Web3 故障不能阻断玩家进入对局。")

    add_heading(doc, "1. 文档目标与范围", 1)
    add_body(doc, "本文把既有链上资产能力转换为游戏侧可实现、可测试的需求。它规定 Unity 客户端、权威游戏服务器与游戏后端需要提供的体验和接口，但不重复定义 Solidity 合约内部实现。")
    add_list_item(doc, "范围内：玩家账号与钱包绑定、皮肤库存与装备、奖励领取、电竞赛事、对局结果存证触发、交易入口、失败降级和验收。", bullet_id)
    add_list_item(doc, "范围外：移动/射击/碰撞算法、武器数值平衡、地图设计、反作弊算法细节、合约代码审计与主网运维。", bullet_id)
    add_list_item(doc, "术语约定：本文中的“赛事”指 TournamentEscrow 对应的电竞赛事奖池托管，不指 NFT 竞价拍卖。", bullet_id)

    add_heading(doc, "2. 产品定位与体验目标", 1)
    add_body(doc, "产品是一款多人联网 FPS。玩家以传统游戏账号进入游戏，钱包作为账号的可选资产属性；即使玩家从未接触钱包，也应能完成登录、匹配、战斗和基础成长。Web3 只增强所有权、奖励、赛事资金透明度与资产流通。")
    add_table(doc,
        ["目标", "游戏侧表现", "成功信号"],
        [
            ["FPS 流畅优先", "战斗期间不发生网关或链调用", "RPC 中断时对局帧率、输入和结算采集不受影响"],
            ["低门槛使用", "先玩游戏，后绑定钱包；默认皮肤始终可用", "未绑钱包玩家能完成完整普通对局"],
            ["资产真实可拥有", "已确认 NFT 可装备、展示并跳转交易", "转移后库存与可装备状态在刷新周期内正确变化"],
            ["结果可验证", "对局结果由服务器生成标准结果包并提交存证任务", "同一结果包可独立复算出相同哈希"],
            ["赛事资金透明", "报名费、奖池、状态、名次与领奖进度可见", "取消、超时、结算三条路径均能解释且可验收"],
        ],
        [1800, 3960, 3600],
    )

    add_heading(doc, "3. 系统边界与职责", 1)
    add_callout(doc, "边界", "Unity 负责体验与表现；权威游戏服务器负责可信战斗结果；游戏后端负责账号、缓存、幂等与风险控制；Web3 资产服务负责签名、交易提交和链上读取；合约只负责最终所有权、托管与不可覆盖的结果承诺。")
    add_table(doc,
        ["组件", "必须负责", "明确禁止"],
        [
            ["Unity 客户端", "大厅 UI、衣柜、装备选择、绑定跳转、奖励/赛事状态展示、失败提示", "保存私钥；直连链；在战斗中调用资产网关；信任本地资产声明"],
            ["权威游戏服务器", "移动/命中/比分权威判定；开局资产核验；生成最终结果包；推送结算", "接受客户端自报 NFT 所有权或最终比分"],
            ["游戏后端", "playerId、会话、钱包绑定、奖励状态、赛事业务状态、结果包 API、重试与审计", "把钱包地址当作唯一游戏账号；无幂等地重复发奖"],
            ["Web3 资产服务", "链读取缓存、SIWE 验签、交易提交、确认追踪、合约事件同步", "影响实时战斗；向 Unity 下发签名私钥"],
            ["链上合约", "NFT 所有权、发行上限、奖励铸造、赛事托管、结果哈希存证", "保存逐帧战斗、等级经验、弹道、聊天或反作弊封禁"],
        ],
        [1600, 4100, 3660],
    )
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("主链路：Unity 大厅 → 游戏后端 → Web3 资产服务 → EVM 合约；战斗链路：Unity 客户端 ↔ 权威游戏服务器（与链完全隔离）。")
    set_run_font(r, size=10.5, bold=True, color=NAVY)

    add_heading(doc, "4. 核心玩家循环", 1)
    for text in [
        "登录传统游戏账号，进入大厅；未绑定钱包时使用默认资产。",
        "在衣柜选择已确认且通过服务器所有权核验的皮肤，保存 loadout。",
        "进入匹配与实时 FPS 对局；对战全程不访问 Web3。",
        "服务器结算比分，生成不可变的标准结果包，并异步创建存证与奖励任务。",
        "玩家回到赛后页/大厅查看战绩、奖励状态和资产同步状态。",
        "玩家可在系统浏览器完成钱包绑定、领取/查看资产或进入 Web 交易市场。",
    ]:
        add_list_item(doc, text, number_id)

    add_heading(doc, "5. 功能需求", 1)
    add_heading(doc, "5.1 账号与钱包绑定", 2)
    reqs = [
        "ACC-001（Must）游戏账号以 playerId 为主身份；钱包地址是可更换的绑定属性，不得替代 playerId。",
        "ACC-002（Must）未绑定钱包的玩家可以登录、匹配、战斗并使用默认皮肤。",
        "ACC-003（Must）Unity 调用 BeginWalletBind 后，使用系统默认浏览器打开 bindUrl；不得使用内嵌 WebView。",
        "ACC-004（Must）绑定页完成 SIWE 签名后，Unity 通过 PollWalletBind 轮询状态：pending、bound、expired、failed。",
        "ACC-005（Must）绑定成功后显示缩写地址并刷新库存；不在 UI、日志或本地文件中保存私钥、助记词或签名密钥。",
        "ACC-006（Should）钱包更换必须经过重新认证，并由后端定义资产迁移/不可迁移提示；Unity 只展示结果。",
    ]
    for x in reqs:
        add_list_item(doc, x, bullet_id, x.split("）")[0] + "）")

    add_heading(doc, "5.2 皮肤库存、装备与展示", 2)
    reqs = [
        "AST-001（Must）衣柜只通过 IGameAssetGateway.GetPlayerAssetsAsync 读取玩家资产，不直接读取合约。",
        "AST-002（Must）所有 tokenId 在 C#、JSON、日志和存档中均按十进制字符串处理，禁止转换为 ulong。",
        "AST-003（Must）库存条目至少包含 tokenId、skinDefId、名称、稀有度、wear、seasonId、缩略图/资源键、状态和 contentHash。",
        "AST-004（Must）只有 confirmed 状态资产可进入正式 loadout；pending 状态可展示“链上确认中”，但不可装备进入对局。",
        "AST-005（Must）开局前由游戏服务器调用 entitlement-check 核验 loadout 所有权，并返回 snapshotId；Unity 无权调用该内部端点。",
        "AST-006（Must）核验失败、资产已转出或 contentHash 不匹配时自动替换为默认皮肤，不阻断开局。",
        "AST-007（Must）其他玩家看到的外观由游戏服务器下发，绝不采信该玩家客户端自报的 tokenId/skinDefId。",
        "AST-008（Should）返回大厅、完成交易或收到资产变更通知后可手动刷新；缓存年龄通过 stalenessSeconds 展示为“同步中”。",
    ]
    for x in reqs:
        add_list_item(doc, x, bullet_id, x.split("）")[0] + "）")

    add_heading(doc, "5.3 实时对战与网络", 2)
    reqs = [
        "COM-001（Must）移动、射击、碰撞、伤害、死亡、复活、比分和胜负均由 Unity 与权威游戏服务器处理。",
        "COM-002（Must）对局场景中禁止调用任何 IGameAssetGateway 方法；资产快照在开局前确定并冻结到本局结束。",
        "COM-003（Must）客户端断线重连不得重新依赖链；服务器依据 snapshotId 和本局会话恢复外观与战斗状态。",
        "COM-004（Must）NFT 皮肤仅改变视觉/音效表现，不提供伤害、射速、命中盒、后坐力或匹配优势。",
        "COM-005（Should）桌面 MVP 以稳定 60 FPS 为最低体验目标；Web3 后台请求不得占用战斗主线程或产生可感知卡顿。",
    ]
    for x in reqs:
        add_list_item(doc, x, bullet_id, x.split("）")[0] + "）")

    add_heading(doc, "5.4 对局结果与链上存证", 2)
    reqs = [
        "MAT-001（Must）只有权威服务器能宣告对局结束并生成 matchId；客户端结果页仅消费服务器结果。",
        "MAT-002（Must）服务器生成 versioned MatchResult 结果包；同一 matchId 的结果包发布后不得在原记录上静默修改。",
        "MAT-003（Must）结果包按 UTF-8 编码，使用 RFC 8785 JSON Canonicalization Scheme 规范化，再计算 keccak256；任何语言实现必须得到同一 resultHash。",
        "MAT-004（Must）matchId 的链上键为 keccak256(UTF-8(matchId))；空 matchId 或空 resultHash 不得进入提交队列。",
        "MAT-005（Must）游戏后端保存完整结果包并提供 GET /v1/matches/{matchId}；链上只写 resultHash。",
        "MAT-006（Must）存证任务异步执行并支持批量；失败仅影响“已验证”标记，不影响赛后页、排行或下一局。",
        "MAT-007（Should）赛后页展示 pending、attested、failed 三种存证状态，并提供区块浏览器/验证页入口。",
    ]
    for x in reqs:
        add_list_item(doc, x, bullet_id, x.split("）")[0] + "）")

    add_heading(doc, "5.5 奖励与 NFT 发放", 2)
    reqs = [
        "RWD-001（Must）奖励先归属于 playerId；只有已绑定钱包且通过反作弊/风控闸门后才能进入铸造。",
        "RWD-002（Must）MVP 默认采用后端 mintDirect（push）路径，玩家不承担 gas；高价值奖励可切换为 voucher（pull）路径。",
        "RWD-003（Must）每个奖励使用稳定 rewardId；合约 requestId 必须由 matchId + playerId + rewardSlot 唯一派生，前后端都要幂等。",
        "RWD-004（Must）Unity 的奖励状态包括 earned、held、claimable、processing、pending_chain、confirmed、failed。",
        "RWD-005（Must）重复点击、客户端重试、网络超时或服务器重复推送不得造成重复 NFT。",
        "RWD-006（Must）铸造失败可重试；界面不得在链确认前把奖励标为已可装备。",
        "RWD-007（Should）发奖成功后自动刷新衣柜，并在奖励卡片展示 tokenId、稀有度和查看资产入口。",
    ]
    for x in reqs:
        add_list_item(doc, x, bullet_id, x.split("）")[0] + "）")

    add_heading(doc, "5.6 电竞赛事与奖池", 2)
    reqs = [
        "TRN-001（Must）游戏内提供赛事列表、详情、报名确认、参赛状态、奖池、规则、报名/结果截止时间和分奖比例。",
        "TRN-002（Must）创建信息中明确显示 organizer、resultSubmitter、entryFee、人数上下限、organizerFeeBps 与 payoutBps；报名后不得由游戏侧暗改。",
        "TRN-003（Must）报名支付、赞助、领奖与退款属于链上交易，Unity 只展示意图和状态，并跳转系统浏览器完成签名/支付。",
        "TRN-004（Must）赛事状态与合约保持一致：Open、Settled、Cancelled；Settled 和 Cancelled 为终态。",
        "TRN-005（Must）人数不足、组织者在允许时间主动取消、结果提交超时三类取消原因必须分别展示；Cancelled 后提供 claimRefund 入口。",
        "TRN-006（Must）结算 winners 顺序与 payoutBps 名次一一对应，且 resultHash 必须引用同一赛事最终对局结果包。",
        "TRN-007（Must）奖金采用玩家主动领取；某一领奖账户失败不得影响其他玩家。Unity 分账户展示可领取金额。",
        "TRN-008（Should）所有金额同时显示 MON 原始单位的格式化值与网络标识，避免把测试网资产误认为真实货币。",
    ]
    for x in reqs:
        add_list_item(doc, x, bullet_id, x.split("）")[0] + "）")

    add_heading(doc, "5.7 交易入口", 2)
    reqs = [
        "MKT-001（Must）MVP 游戏内不实现托管钱包和完整交易签名；从资产详情跳转系统浏览器进入 Web 市场。",
        "MKT-002（Must）返回游戏后刷新库存；若 NFT 已转出，当前 loadout 在下一次开局核验时回退默认皮肤。",
        "MKT-003（Should）资产详情可展示可转让提示、版税提示与外部交易风险说明，不承诺版税一定被所有市场执行。",
    ]
    for x in reqs:
        add_list_item(doc, x, bullet_id, x.split("）")[0] + "）")

    add_heading(doc, "6. Unity 客户端模块需求", 1)
    add_table(doc,
        ["模块", "输入", "输出/行为", "禁止依赖"],
        [
            ["GameAssetService", "IGameAssetGateway、会话 Token", "统一超时、取消、错误映射和缓存；仅大厅调用", "区块链 SDK、私钥、战斗对象"],
            ["ArmoryController", "PlayerAssets、资源目录", "筛选 confirmed、预览、装备、默认回退", "客户端自认所有权"],
            ["WalletBindController", "绑定会话与 bindUrl", "系统浏览器跳转、轮询、过期重试", "内嵌 WebView、助记词输入"],
            ["RewardController", "rewardId、RewardStatus", "状态机、重试、成功刷新", "直接调用合约 mint"],
            ["TournamentController", "赛事 REST 模型", "列表/详情/报名意图/状态刷新/领奖入口", "本地计算最终奖金额"],
            ["MatchResultPresenter", "服务器 MatchResult", "赛后页与存证状态展示", "客户端生成最终比分"],
            ["Web3DegradePolicy", "错误类别、缓存年龄", "默认皮肤、功能灰显、用户可理解提示", "退出游戏或阻断匹配"],
        ],
        [1700, 2150, 3310, 2200],
    )

    add_heading(doc, "7. 对外接口与数据契约", 1)
    add_heading(doc, "7.1 Unity 网关接口", 2)
    add_table(doc,
        ["方法", "调用场景", "关键返回", "失败策略"],
        [
            ["GetPlayerAssetsAsync", "大厅/衣柜/交易返回", "PlayerAssets + stalenessSeconds", "缓存可用则展示；否则默认皮肤"],
            ["BeginWalletBindAsync", "玩家点击绑定", "sessionId + bindUrl + expiresAt", "允许重试，不影响游戏"],
            ["PollWalletBindAsync", "绑定浏览器已打开", "pending/bound/expired/failed", "到期停止轮询并提示重开"],
            ["RequestClaimAsync", "玩家领取奖励", "ClaimTicket / 当前状态", "同 rewardId 幂等返回"],
            ["PollRewardAsync", "领奖处理中", "RewardStatus + tokenId/txHash", "退避轮询，可跨场景恢复"],
        ],
        [2400, 2250, 2570, 2140],
    )
    add_body(doc, "实现约束：开发期使用 MockGameAssetGateway；联调期替换 HttpGameAssetGateway。上层 UI 与业务控制器不得因实现切换而修改。所有异步调用接受 CancellationToken，并在场景销毁时取消。")

    add_heading(doc, "7.2 MatchResult 结果包（游戏服务器输出）", 2)
    add_table(doc,
        ["字段", "类型", "必填", "规则"],
        [
            ["version", "string", "是", "固定为 schema 版本，如 1.0"],
            ["matchId", "string", "是", "全局唯一；生成后不可复用"],
            ["modeId / mapId", "string", "是", "稳定配置 ID，不使用本地化名称"],
            ["startedAt / endedAt", "int64", "是", "UTC Unix 秒；endedAt ≥ startedAt"],
            ["serverBuild", "string", "是", "可追溯到服务端构建版本"],
            ["tournamentId", "string/null", "否", "赛事局对应链上 tournamentId"],
            ["players", "array", "是", "按 playerId 升序；至少含 playerId、teamId、kills、deaths、score、placement、result"],
            ["antiCheatState", "string", "是", "passed/held/rejected；held/rejected 禁止立即发高价值奖励"],
            ["rewardSlots", "array", "否", "slot、playerId、rewardId；作为幂等派生输入"],
        ],
        [1900, 1500, 900, 5060],
    )

    add_heading(doc, "7.3 服务器专用资产核验", 2)
    add_body(doc, "POST /internal/v1/entitlement-check 只能由游戏服务器调用。请求包含 playerId、wallet、待装备 tokenId 列表和 matchId；返回 allowed、snapshotId、resolvedSkins、rejectedTokenIds 与 cacheAge。若依赖不可用，服务端按策略返回默认皮肤快照，而不是拒绝加入对局。")

    add_heading(doc, "8. 界面与交互清单", 1)
    add_table(doc,
        ["界面", "主要内容", "关键操作", "异常态"],
        [
            ["大厅", "账号、钱包摘要、奖励/赛事入口", "进入匹配、刷新资产", "Web3 不可用横幅但匹配仍可用"],
            ["衣柜", "默认皮肤、confirmed/pending NFT", "预览、装备、查看资产", "同步中、已转出、资源哈希失败"],
            ["钱包绑定", "流程说明、二维码/链接状态", "打开系统浏览器、重新绑定", "过期、拒签、地址冲突"],
            ["赛后结果", "比分、名次、奖励、存证状态", "查看奖励、验证结果", "存证失败不遮挡战绩"],
            ["奖励中心", "earned/held/processing/confirmed", "领取、重试、查看 NFT", "重复请求、超时、链确认中"],
            ["赛事列表/详情", "奖池、人数、费用、截止、提交方、分奖", "报名、赞助、查看状态", "报名关闭、已满、余额不足"],
            ["赛事结算", "名次、resultHash、可领奖/退款金额", "跳转领取/退款", "提交超时、取消原因、交易失败"],
            ["资产详情", "tokenId、款式、稀有度、磨损、赛季", "打开 Web 市场/浏览器", "网络错误、资产已转移"],
        ],
        [1650, 3050, 2360, 2300],
    )

    add_heading(doc, "9. 非功能与安全需求", 1)
    nfrs = [
        "NFR-001 性能：网关响应、JSON 解析、图片加载与缓存更新不得在战斗主线程同步阻塞；战斗场景不启动资产轮询。",
        "NFR-002 可用性：RPC、资产后端或缓存故障时，普通对局使用默认皮肤继续；只有资产、领奖、报名/交易入口降级。",
        "NFR-003 安全：客户端和客户端日志不得包含私钥、助记词、服务签名密钥、管理员地址凭据或后端内部令牌。",
        "NFR-004 信任：客户端上报的 tokenId、比分、名次、奖励资格和交易状态均不作为最终依据。",
        "NFR-005 幂等：绑定会话、结果推送、奖励请求、链交易提交和事件消费均可安全重试。",
        "NFR-006 可观测：记录 requestId、matchId、rewardId、playerId、tournamentId、txHash 和错误码；禁止记录签名原文与敏感认证材料。",
        "NFR-007 最终性：链上事件未达到 Web3 服务设定的最终确认策略前，资产/奖励只能显示 pending，不得用于正式 loadout。",
        "NFR-008 资源完整性：contentHash 校验失败时使用默认资源并上报告警；不得加载未经验证的远程可执行内容。",
        "NFR-009 可访问性：关键状态不只依赖颜色；pending、failed、confirmed 同时使用文字和图标。",
        "NFR-010 法务提示：NFT 所有权不等于游戏账号所有权；游戏内可用性仍受服务运营、反作弊与内容规则约束。",
    ]
    for x in nfrs:
        add_list_item(doc, x, bullet_id)

    add_heading(doc, "10. 失败降级矩阵", 1)
    add_table(doc,
        ["故障", "玩家可继续做什么", "受限功能", "游戏侧处理"],
        [
            ["链 RPC 不可用", "登录、匹配、战斗、查看缓存库存", "刷新资产、领奖、报名/交易", "显示维护提示；使用缓存或默认皮肤"],
            ["资产后端超时", "匹配与普通对局", "绑定、库存刷新、奖励", "取消请求、可重试、禁止无限 loading"],
            ["缓存与数据库同时不可用", "使用默认皮肤进入对局", "NFT 装备", "服务器生成默认 snapshotId"],
            ["contentHash 不一致", "继续游戏", "对应 NFT 外观", "替换默认资源并告警"],
            ["奖励交易失败", "继续下一局", "该奖励暂不可用", "保持 failed/processing 状态并幂等重试"],
            ["结果存证失败", "查看赛后结果、继续游戏", "“已验证”标识", "异步重试；禁止修改已发布结果包"],
            ["赛事提交方超时", "查看赛事与退款资格", "正常结算", "到 resultDeadline 后展示取消/退款入口"],
        ],
        [1900, 2600, 2100, 2760],
    )

    add_heading(doc, "11. MVP 优先级与明确不做", 1)
    add_heading(doc, "11.1 MVP 必须完成", 2)
    for x in [
        "Mock 与 HTTP 网关可无缝切换；大厅、衣柜和默认皮肤降级闭环。",
        "传统账号登录 + 系统浏览器钱包绑定 + 库存读取。",
        "开局服务器 entitlement-check + 本局资产 snapshot。",
        "对局结果包、哈希规范、异步存证状态和赛后页。",
        "奖励 mintDirect 幂等链路与 confirmed 后可装备。",
        "赛事浏览、报名跳转、Open/Settled/Cancelled 展示、领奖/退款入口。",
        "外部 Web 市场入口与交易返回后的库存刷新。",
    ]:
        add_list_item(doc, x, bullet_id)
    add_heading(doc, "11.2 本版本明确不做", 2)
    for x in [
        "逐帧战斗、每枪弹道、经验、等级、匹配分或封禁记录上链。",
        "游戏内 ERC-20、质押、挖矿、收益玩法、竞猜或押注。",
        "宝箱/开箱及任何需要随机数且可能触发博彩合规的问题。",
        "Unity 内嵌钱包、私钥托管、内嵌 WebView 签名或直接 RPC。",
        "游戏内完整 NFT 拍卖市场；MVP 使用外部 Web 市场。",
        "NFT 给予战斗数值优势或付费获胜能力。",
    ]:
        add_list_item(doc, x, bullet_id)

    add_heading(doc, "12. 验收测试", 1)
    add_table(doc,
        ["编号", "场景", "验收标准"],
        [
            ["AC-01", "未绑定钱包进入游戏", "可登录、匹配、使用默认皮肤完成一局；无阻断弹窗"],
            ["AC-02", "绑定钱包", "系统浏览器完成签名；Unity 从 pending 进入 bound 并刷新库存"],
            ["AC-03", "持有 NFT 装备", "confirmed 资产可选；服务器核验通过并向所有客户端下发正确外观"],
            ["AC-04", "伪造 tokenId", "客户端改包不生效；服务器替换默认皮肤并记录拒绝原因"],
            ["AC-05", "战斗中 RPC 宕机", "移动、射击、命中和结算采集无异常；只影响赛后 Web3 状态"],
            ["AC-06", "同一奖励重复请求", "多次点击/重试只产生一个 requestId 对应的一件 NFT"],
            ["AC-07", "链上奖励等待确认", "状态显示 pending_chain；确认前不可装备，确认后自动/手动刷新可见"],
            ["AC-08", "结果哈希互操作", "C# 与后端实现对同一标准结果包计算出相同 keccak256"],
            ["AC-09", "同一 matchId 二次存证", "第二次提交被拒绝；游戏侧保留原结果并上报告警"],
            ["AC-10", "赛事正常结算", "名次与 payoutBps 对齐；各获奖者独立看到可领取金额"],
            ["AC-11", "赛事人数不足/超时", "显示准确取消原因；每位付款方可独立进入退款流程"],
            ["AC-12", "资产交易后返回", "库存刷新；已转出 NFT 不再可装备，既有 loadout 下局回退默认"],
            ["AC-13", "资源哈希错误", "不崩溃、不拦截匹配；使用默认外观并产生日志/告警"],
            ["AC-14", "Mock 切 HTTP", "仅替换网关实现与配置；上层 UI/控制器不改代码"],
        ],
        [1000, 2800, 5560],
    )

    add_heading(doc, "13. 交付依赖与待游戏团队锁定项", 1)
    add_body(doc, "下列项目不由链上能力决定，需在游戏立项阶段单独锁定；它们不会改变本文的游戏/Web3 分层原则。")
    add_table(doc,
        ["待锁定项", "建议责任人", "对本文的影响"],
        [
            ["Unity LTS 具体版本、渲染管线与目标平台", "技术负责人", "影响资源包、性能基线和 SDK 打包方式"],
            ["联机方案、服务器 tick 与部署拓扑", "网络/服务器负责人", "影响 MatchResult 生产点与重连策略"],
            ["首发模式、人数、地图与计分规则", "主策划", "影响结果包字段、名次和赛事规则"],
            ["皮肤资源目录与 contentHash 生成流水线", "美术技术/构建负责人", "影响资源校验与版本兼容"],
            ["奖励经济、发行量与稀有度", "经济/运营策划", "受链上 maxSupply 约束；上链后只能降不能升"],
            ["高价值奖励反作弊 held 规则", "安全/运营", "决定何时允许 mint，不影响对局结算"],
        ],
        [3100, 1900, 4360],
    )

    add_heading(doc, "14. 链上能力到游戏需求的追踪", 1)
    add_table(doc,
        ["链上模块", "提供的能力", "对应游戏需求"],
        [
            ["GameAssetRegistry", "款式定义、发行上限、contentHash", "AST-003/006、NFR-008、资源目录与稀缺性展示"],
            ["WeaponSkin", "ERC-721 所有权、磨损、赛季、tokenId 编码", "衣柜、装备、资产详情、交易后刷新"],
            ["RewardDistributor", "voucher 与 mintDirect、nonce/requestId 幂等", "RWD-001~007、奖励状态机"],
            ["SkinMarket", "固定价转让与版税信息", "MKT-001~003；MVP 仅外部 Web 入口"],
            ["MatchAttestation", "不可覆盖 resultHash、批量存证", "MAT-001~007、赛后验证状态"],
            ["TournamentEscrow", "报名/赞助托管、分奖、取消与退款", "TRN-001~008、赛事 UI 与状态"],
        ],
        [2100, 3290, 3970],
    )

    add_heading(doc, "附录 A：来源与假设", 1)
    add_body(doc, "来源：本文依据 timothyshen/web3-fps-assets 仓库中现有合约、接口、Unity SDK 与集成文档整理。")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    add_hyperlink(p, "timothyshen/web3-fps-assets（GitHub）", "https://github.com/timothyshen/web3-fps-assets")
    for x in [
        "首个版本采用 EVM 测试环境；具体网络配置由 Web3 团队提供，Unity 不硬编码 chainId 或合约地址。",
        "MVP 面向桌面多人 FPS；移动端、主机平台和跨平台账号合并不在本文范围。",
        "奖励默认使用后端直铸；voucher 路径保留给高价值奖励或未来 gas 代付方案。",
        "赛事采用电竞奖池托管语义，不包含 NFT 竞拍。",
        "所有具体武器数值、地图、模式人数与匹配规则由独立 GDD 决定。",
    ]:
        add_list_item(doc, x, bullet_id)

    # Normalize all table typography and prevent row splitting where practical.
    for table in doc.tables:
        for row in table.rows:
            tr_pr = row._tr.get_or_add_trPr()
            cant_split = OxmlElement("w:cantSplit")
            tr_pr.append(cant_split)
            for cell in row.cells:
                for p in cell.paragraphs:
                    for run in p.runs:
                        if run.font.size is None:
                            set_run_font(run, size=9.2)

    # Keep headings with following content and avoid orphaned body lines.
    for p in doc.paragraphs:
        if p.style.name.startswith("Heading"):
            p.paragraph_format.keep_with_next = True
        p.paragraph_format.widow_control = True

    doc.core_properties.title = "Unity 3D Web3 FPS 游戏设计需求文档"
    doc.core_properties.subject = "游戏侧与链上资产层分离的需求基线"
    doc.core_properties.author = "Codex"
    doc.core_properties.keywords = "Unity 3D, FPS, Web3, NFT, Tournament, Match Attestation, PRD"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build_document()
